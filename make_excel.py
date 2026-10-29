"""midas Gen NX の MGT/MGTX（モデル）と ANL（解析結果）から、計算書の節ごとに応力を確認できる Excel を作る。

  py -3 make_excel.py                         … ファイル選択の画面（start.bat から）
  py -3 make_excel.py model.mgtx result.anl -o out.xlsx [--face-beam 0] [--face-col 0]

シート構成（概要 → 4-1 RC柱 → 4-2 RC梁 → 4-3 耐震壁 → 4-4 支点まわり → 断面別まとめ → グループ別まとめ）。
RC のシートは「荷重ケースごとの応力」と「長期／短期／終局のまとめ（数式）」を同じ表に置く。
ANL には梁要素の断面力しか出ていないため、板要素（壁）の応力と支点反力は出せない（概要シートに明記）。
"""
import argparse
import collections
import math
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BS = chr(92)
PTS = ['I', '1/4', 'CNT', '3/4', 'J']
COMPS = ['N', 'Qy', 'Qz', 'T', 'My', 'Mz']
COMP_UNIT = {'N': 'kN', 'Qy': 'kN', 'Qz': 'kN', 'T': 'kN·m', 'My': 'kN·m', 'Mz': 'kN·m'}
COMP_NAME = {'N': '軸力 N', 'Qy': 'せん断 Qy', 'Qz': 'せん断 Qz', 'T': 'ねじり T', 'My': '曲げ My', 'Mz': '曲げ Mz'}
SIGMA = '?SIGMA?'
FLOAT = re.compile(r'^[-+]?(?:\d+\.?\d*|\.\d+)(?:[Ee][-+]?\d+)?$')


def say(msg):
    print(msg, flush=True)


# ───────────────────────── MGT / MGTX ─────────────────────────

def expand_ids(text):
    ids = []
    for tok in str(text).split():
        m = re.match(r'^(\d+)to(\d+)(?:by(\d+))?$', tok, re.I)
        if m:
            a, b, step = int(m[1]), int(m[2]), int(m[3] or 1)
            if step >= 1 and abs(b - a) <= 200000:
                ids.extend(range(a, b + 1, step) if a <= b else range(a, b - 1, -step))
        elif tok.isdigit():
            ids.append(int(tok))
    return ids


def read_mgt(path):
    text = open(path, 'rb').read().decode('cp932', 'replace')
    secs = {}
    key = None
    for line in text.split('\n'):
        m = re.match(r'\s*\*([A-Z0-9_-]+)', line)
        if m:
            key = m.group(1).upper()
            secs.setdefault(key, [])
            continue
        if key:
            secs[key].append(line.rstrip('\r'))

    def logical(name, join_backslash=True):
        out, cur = [], ''
        for raw in secs.get(name, []):
            s = raw.strip()
            if not s or s[0] in ';$*':
                continue
            cur += s
            if join_backslash and cur.endswith(BS):
                cur = cur[:-1] + ' '
                continue
            out.append(cur)
            cur = ''
        return out

    unit = [x.strip().upper() for x in (logical('UNIT') or ['KN, MM'])[0].split(',')]
    length_to_mm = {'MM': 1.0, 'CM': 10.0, 'M': 1000.0}.get(unit[1] if len(unit) > 1 else 'MM', 1.0)

    nodes = {}
    for l in secs.get('NODE', []):
        p = [x.strip() for x in l.split(',')]
        if p and p[0].isdigit() and len(p) >= 4:
            nodes[int(p[0])] = tuple(float(x) * length_to_mm for x in p[1:4])
    elements = {}
    for l in secs.get('ELEMENT', []):
        p = [x.strip() for x in l.split(',')]
        if len(p) >= 6 and p[0].isdigit():
            elements[int(p[0])] = dict(id=int(p[0]), type=p[1].upper(), mat=p[2], sec=int(p[3]) if p[3].isdigit() else 0,
                                       nodes=[int(x) for x in p[4:8] if x.isdigit() and x != '0'])
    sections = {}
    for l in secs.get('SECTION', []):
        p = [x.strip() for x in l.split(',')]
        if len(p) >= 3 and p[0].isdigit():
            sections[int(p[0])] = p[2]
    thickness = {}
    for l in secs.get('THICKNESS', []):
        p = [x.strip() for x in l.split(',')]
        if len(p) >= 5 and p[0].isdigit():
            try:
                thickness[int(p[0])] = float(p[4]) * length_to_mm
            except ValueError:
                pass
    groups = collections.OrderedDict()
    for l in logical('GROUP'):
        p = [x.strip() for x in l.split(',')]
        if len(p) >= 3 and p[0]:
            groups[p[0]] = expand_ids(p[2])
    supports = set()
    for l in logical('CONSTRAINT'):
        p = [x.strip() for x in l.split(',')]
        if len(p) >= 2 and re.match(r'^[01]{6,7}$', p[1]):
            supports.update(expand_ids(p[0]))
    return dict(nodes=nodes, elements=elements, sections=sections, thickness=thickness, groups=groups, supports=supports,
                grids=read_grids(secs, length_to_mm))


# ───────────────────────── 通り芯（midas の寸法線・テキスト） ─────────────────────────

def dimension_point(view, pt):
    rx, ry, rz = (math.radians(a) for a in view['rotation'])
    x, y, z = (v * view['scale'] for v in pt)
    y1, z1 = y * math.cos(rx) - z * math.sin(rx), y * math.sin(rx) + z * math.cos(rx)
    x2, z2 = x * math.cos(ry) + z1 * math.sin(ry), -x * math.sin(ry) + z1 * math.cos(ry)
    x3, y3 = x2 * math.cos(rz) - y1 * math.sin(rz), x2 * math.sin(rz) + y1 * math.cos(rz)
    return (x3 + view['origin'][0], y3 + view['origin'][1], z2 + view['origin'][2])


def read_grids(secs, length_to_mm):
    """*DIMENSION（原点）と *DIMENSION-LINETEXT（線・文字）から、X 系・Y 系の通り芯を作る。

    符号は文字に一番近い線の端に付ける。名前が X／BX で始まれば X 系、Y／BY なら Y 系。
    """
    views = {}
    for l in secs.get('DIMENSION', []):
        p = [x.strip() for x in l.split(',')]
        if len(p) >= 8 and p[0] and not p[0].startswith(';'):
            try:
                v = [float(x) for x in p[1:8]]
            except ValueError:
                continue
            views[p[0]] = dict(name=p[0], scale=v[0], origin=tuple(c * length_to_mm for c in v[1:4]), rotation=v[4:7], lines=[], texts=[])
    for l in secs.get('DIMENSION-LINETEXT', []):
        p = [x.strip() for x in l.split(',')]
        if len(p) < 11 or p[0] not in views:
            continue
        try:
            kind = int(p[1])
            co = [float(x) * length_to_mm for x in p[5:11]]
        except ValueError:
            continue
        view = views[p[0]]
        a, b = dimension_point(view, co[:3]), dimension_point(view, co[3:])
        if kind == 0:
            view['lines'].append((a, b))
        elif kind == 1 and p[2]:
            view['texts'].append((p[2], a))
    grids = []
    for view in views.values():
        named = []
        for name, at in view['texts']:
            best = None
            for line in view['lines']:
                for end in (0, 1):
                    d = math.hypot(line[end][0] - at[0], line[end][1] - at[1])
                    if best is None or d < best[0]:
                        best = (d, line)
            if best and best[0] < 5000:
                named.append((name, best[1]))
        for name, line in named:
            m = re.match(r'^B?([XY])', name)
            if m and math.hypot(line[1][0] - line[0][0], line[1][1] - line[0][1]) > 1:
                grids.append(dict(view=view['name'], name=name, family=m.group(1), p0=line[0][:2], p1=line[1][:2], z=view['origin'][2]))
    return grids


class GridIndex:
    """平面上の点・部材が、midas の通り芯のどこにあるかを、計算書のチャートと同じ書き方で文字にする。

    点：「BX2・BY1」（通り上）、「BX2+1250・BY2」（通りから 1250 mm 離れている。＋は番号が進む側）
    部材：通りに沿う部材は「BY1=BX1～BX2」（通り＝始点～終点）、沿わない部材は「始点 ～ 終点」。
    """

    AXIS = {'X': (1.0, 0.0), 'Y': (0.0, 1.0)}   # X 系は西→東、Y 系は南→北の順に並べる

    def __init__(self, grids, tol=100.0):
        self.tol = tol
        self.views = collections.OrderedDict()
        for g in grids:
            self.views.setdefault(g['view'], dict(z=g['z'], fam={}))['fam'].setdefault(g['family'], []).append(g)
        for v in self.views.values():
            for fam, lines in v['fam'].items():
                ax = self.AXIS[fam]
                for g in lines:
                    d = (g['p1'][0] - g['p0'][0], g['p1'][1] - g['p0'][1])
                    length = math.hypot(*d)
                    n = (-d[1] / length, d[0] / length)
                    if n[0] * ax[0] + n[1] * ax[1] < 0:
                        n = (-n[0], -n[1])
                    g['n'], g['d'] = n, (d[0] / length, d[1] / length)
                    g['mid'] = ((g['p0'][0] + g['p1'][0]) / 2, (g['p0'][1] + g['p1'][1]) / 2)
                lines.sort(key=lambda g: g['mid'][0] * ax[0] + g['mid'][1] * ax[1])

    @property
    def available(self):
        return bool(self.views)

    def view_for(self, ref_z):
        return min(self.views.values(), key=lambda v: abs(v['z'] - ref_z)) if self.views else None

    @staticmethod
    def _offset(g, pt):
        """通りから点までの距離（mm）。＋は番号が進む側。"""
        return g['n'][0] * (pt[0] - g['p0'][0]) + g['n'][1] * (pt[1] - g['p0'][1])

    def _pos(self, lines, pt):
        k = min(range(len(lines)), key=lambda i: abs(self._offset(lines[i], pt)))
        o = self._offset(lines[k], pt)
        return lines[k]['name'] if abs(o) <= self.tol else f'{lines[k]["name"]}{o:+.0f}'

    def label(self, pt, ref_z):
        """pt: 全体座標の点、ref_z: 通りの面を選ぶための高さ（部材の最低点）。"""
        view = self.view_for(ref_z)
        if not view:
            return ''
        return '・'.join(self._pos(view['fam'][f], pt) for f in ('X', 'Y') if f in view['fam'])

    def member(self, a, b, ref_z):
        """部材（始点 a・終点 b）の通り表記。通りに沿っていれば「通り＝始点～終点」。"""
        view = self.view_for(ref_z)
        if not view:
            return ''
        u = (b[0] - a[0], b[1] - a[1])
        length = math.hypot(*u)
        if length > 1:
            u = (u[0] / length, u[1] / length)
            for fam, other in (('X', 'Y'), ('Y', 'X')):
                lines = view['fam'].get(fam)
                if not lines or other not in view['fam']:
                    continue
                for g in lines:
                    if abs(u[0] * g['d'][1] - u[1] * g['d'][0]) < 0.05 and abs(self._offset(g, a)) <= self.tol and abs(self._offset(g, b)) <= self.tol:
                        return f'{g["name"]}={self._pos(view["fam"][other], a)}～{self._pos(view["fam"][other], b)}'
        return f'{self.label(a, ref_z)} ～ {self.label(b, ref_z)}'


def point_on_element(model, element, pt):
    nodes = model['nodes']
    a, b = nodes[element['nodes'][0]], nodes[element['nodes'][1]]
    t = {'I': 0.0, '1/4': 0.25, 'CNT': 0.5, '3/4': 0.75, 'J': 1.0}[pt]
    return tuple(a[i] + (b[i] - a[i]) * t for i in range(3))


def element_grid(model, gi, eid, pt):
    e = model['elements'][eid]
    ref_z = min(model['nodes'][n][2] for n in e['nodes'][:2])
    return gi.label(point_on_element(model, e, pt), ref_z)


# ───────────────────────── ANL ─────────────────────────

def read_anl(path):
    data = open(path, 'rb').read()
    # GEN NX は表ごとに文字コードが違い、荷重ケース名の「Σ」が CP932(83 B0)／CP949(A5 D2) で出てくる。
    data = data.replace(b'\x83\xb0', SIGMA.encode()).replace(b'\xa5\xd2', SIGMA.encode())
    raw_lines = data.split(b'\n')
    del data
    dec = lambda b: b.decode('cp932', 'replace').replace(SIGMA, 'Σ')
    lines = [dec(b) for b in raw_lines[:3000]]

    m_unit = None
    for l in lines[:600]:
        m_unit = re.search(r':\s*(kN|N|tonf|kgf)\s*,\s*(mm|cm|m)\b', l)
        if m_unit:
            break
    force_unit, length_unit = (m_unit.group(1), m_unit.group(2)) if m_unit else ('kN', 'mm')
    f_fac = {'kN': 1.0, 'N': 0.001, 'tonf': 9.80665, 'kgf': 0.00980665}[force_unit]
    l_fac = {'mm': 0.001, 'cm': 0.01, 'm': 1.0}[length_unit]
    scale = [f_fac] * 3 + [f_fac * l_fac] * 3

    # 短縮名（NΣDL+S~4）→ 正式名称（要約表）
    full = {}
    for i, l in enumerate(lines[:3000]):
        if '要約表' in l and '<<' in l:
            rules = 0
            for l2 in lines[i + 1:i + 200]:
                s = l2.strip()
                if s.startswith('-----'):
                    rules += 1
                    if rules == 2:
                        break
                    continue
                if rules == 1 and s:
                    t = s.split()
                    if len(t) >= 2:
                        full[t[0]] = t[1]
            break

    results = collections.defaultdict(lambda: collections.defaultdict(dict))
    case_order = []
    seen = set()
    cur_elem = cur_case = None
    has_data = False
    sec_of = {}
    for raw in raw_lines:
        t = raw.split()
        if not t:
            continue
        t = [dec(x) for x in t[:12]]
        if has_data and t[0].startswith('_' * 10):
            break
        if len(t) >= 11 and t[0].isdigit() and t[1].isdigit() and t[2].isdigit() and t[4] in PTS and all(FLOAT.match(x) for x in t[5:11]):
            cur_elem, cur_case, pt, vals = int(t[0]), t[3], t[4], t[5:11]
            sec_of[cur_elem] = int(t[2])
        elif cur_elem is not None and len(t) >= 7 and t[0] in PTS and all(FLOAT.match(x) for x in t[1:7]):
            pt, vals = t[0], t[1:7]
        elif cur_elem is not None and len(t) >= 8 and t[1] in PTS and all(FLOAT.match(x) for x in t[2:8]):
            cur_case, pt, vals = t[0], t[1], t[2:8]
        else:
            continue
        has_data = True
        name = full.get(cur_case, cur_case)
        if name not in seen:
            seen.add(name)
            case_order.append(name)
        results[cur_elem][name][pt] = [float(v) * s for v, s in zip(vals, scale)]
    return dict(results=results, cases=case_order, sec_of=sec_of, units=f'{force_unit}, {length_unit}')


def classify(name):
    """荷重ケース名から 長期／短期／終局 を仮決めする（概要シートで直せる）。"""
    if '終局' in name:
        return '終局'
    if '+' not in name and name.rstrip().endswith('DL'):
        return '長期'
    return '短期'


# ───────────────────────── 部材のまとめ ─────────────────────────

class Line:
    """一直線上に並べた梁要素の列（RC梁は連続する要素を 1 つの部材として扱う）。"""

    def __init__(self, elems, model):
        self.elems = elems
        nodes = model['nodes']
        self.node_ids = [elems[0]['nodes'][0]] + [e['nodes'][1] for e in elems]
        self.lengths = [math.dist(nodes[e['nodes'][0]], nodes[e['nodes'][1]]) for e in elems]
        self.total = sum(self.lengths)

    def stations(self, results, case, ci):
        """(距離 s, 値) の並び。要素の 5 点をつなぐ。"""
        out, s0 = [], 0.0
        for e, L in zip(self.elems, self.lengths):
            pts = results.get(e['id'], {}).get(case)
            if not pts or any(p not in pts for p in PTS):
                return None
            for k, p in enumerate(PTS):
                out.append((s0 + L * k / 4, pts[p][ci]))
            s0 += L
        return out

    def point_at(self, model, s):
        """部材の始端から距離 s の点（全体座標）。"""
        nodes = model['nodes']
        s = max(0.0, min(self.total, s))
        s0 = 0.0
        for e, L in zip(self.elems, self.lengths):
            if s <= s0 + L + 1e-9:
                a, b = nodes[e['nodes'][0]], nodes[e['nodes'][1]]
                t = (s - s0) / L if L else 0.0
                return tuple(a[i] + (b[i] - a[i]) * t for i in range(3))
            s0 += L
        return nodes[self.elems[-1]['nodes'][1]]

    def min_z(self, model):
        return min(model['nodes'][n][2] for e in self.elems for n in e['nodes'][:2])

    def value_at(self, results, case, ci, s):
        st = self.stations(results, case, ci)
        if not st:
            return None
        s = max(0.0, min(self.total, s))
        for (a, va), (b, vb) in zip(st, st[1:]):
            if b - a > 1e-9 and a - 1e-9 <= s <= b + 1e-9:
                return va + (vb - va) * (s - a) / (b - a)
        return st[-1][1] if s >= self.total else st[0][1]


def build_chains(elems, model):
    """同じ断面で、他の梁要素・支点に接しない節点でつながる要素を 1 本の部材にする（向きが揃うものだけ）。"""
    frames = [e for e in model['elements'].values() if e['type'] == 'BEAM']
    degree = collections.Counter(n for e in frames for n in e['nodes'][:2])
    by_start = {e['nodes'][0]: e for e in elems}
    by_end = {e['nodes'][1]: e for e in elems}
    taken, chains = set(), []

    def joint(a, b):
        n = a['nodes'][1]
        return b['nodes'][0] == n and degree[n] == 2 and n not in model['supports'] and a['sec'] == b['sec']

    for e in sorted(elems, key=lambda x: x['id']):
        if e['id'] in taken:
            continue
        head = e
        while True:  # 先頭までさかのぼる
            prev = by_end.get(head['nodes'][0])
            if prev and prev['id'] not in taken and joint(prev, head) and prev is not e:
                head = prev
            else:
                break
        chain = [head]
        taken.add(head['id'])
        while True:
            nxt = by_start.get(chain[-1]['nodes'][1])
            if nxt and nxt['id'] not in taken and joint(chain[-1], nxt):
                chain.append(nxt)
                taken.add(nxt['id'])
            else:
                break
        chains.append(Line(chain, model))
    chains.sort(key=lambda c: c.elems[0]['id'])
    return chains


def id_range(elems):
    ids = [e['id'] for e in elems]
    return str(ids[0]) if len(ids) == 1 else f'{ids[0]}〜{ids[-1]}（{len(ids)}本）'


# ───────────────────────── Excel ─────────────────────────

THIN = Side(style='thin', color='B7B7B7')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill('solid', fgColor='D9E2F3')
SUMF = PatternFill('solid', fgColor='FFF2CC')
INPF = PatternFill('solid', fgColor='FFFFCC')
ZEBRA = PatternFill('solid', fgColor='F5F5F5')
BOLD = Font(bold=True)
CLASS_COLOR = {'長期': 'E2EFDA', '短期': 'FCE4D6', '終局': 'EDE1F5'}


def title(ws, text, notes):
    ws['A1'] = text
    ws['A1'].font = Font(bold=True, size=13)
    for i, n in enumerate(notes):
        ws.cell(2 + i, 1, n).font = Font(size=9, color='555555')
    return 2 + len(notes) + 1


def write_case_table(ws, top, info_heads, rows, cases, class_cells):
    """荷重ケース別の応力表と、長期／短期／終局のまとめ（数式）を書く。

    rows: [(info 列の値のリスト, 成分名, [ケースごとの値])]
    class_cells: 概要シートの区分セル参照（ケースの順）
    """
    ni = len(info_heads)
    c_comp, c_unit = ni + 1, ni + 2
    sums = ['長期', '短期 最大', '短期 最小', '短期 支配ケース', '終局 最大', '終局 最小']
    c_sum = c_unit + 1
    c_case = c_sum + len(sums)
    last_case = c_case + len(cases) - 1
    cls_row, head_row, first = top, top + 1, top + 2
    rng = lambda r: f'${get_column_letter(c_case)}{r}:${get_column_letter(last_case)}{r}'
    cls_rng = f'${get_column_letter(c_case)}${cls_row}:${get_column_letter(last_case)}${cls_row}'
    name_rng = f'${get_column_letter(c_case)}${head_row}:${get_column_letter(last_case)}${head_row}'

    ws.cell(cls_row, c_sum - 1, '区分（概要シートで変更）').alignment = Alignment(horizontal='right')
    for j, ref in enumerate(class_cells):
        c = ws.cell(cls_row, c_case + j, f'={ref}')
        c.alignment = Alignment(horizontal='center')
        c.border = BORDER
    heads = list(info_heads) + ['成分', '単位'] + sums + list(cases)
    for j, h in enumerate(heads):
        c = ws.cell(head_row, 1 + j, h)
        c.font, c.fill, c.border = BOLD, (SUMF if c_sum <= 1 + j < c_case else HEAD), BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[head_row].height = 44

    prev, zebra = None, False
    for k, (info, comp, vals) in enumerate(rows):
        r = first + k
        if info != prev:
            zebra = not zebra
            prev = info
        for j, v in enumerate(info):
            ws.cell(r, 1 + j, v)
        ws.cell(r, c_comp, COMP_NAME[comp])
        ws.cell(r, c_unit, COMP_UNIT[comp])
        for j, v in enumerate(vals):
            c = ws.cell(r, c_case + j, None if v is None else round(v, 4))
            c.number_format = '0.0'
        rg = rng(r)
        mx = lambda cls: f'_xlfn.MAXIFS({rg},{cls_rng},"{cls}")'
        mn = lambda cls: f'_xlfn.MINIFS({rg},{cls_rng},"{cls}")'
        L, F, G = get_column_letter(c_sum), get_column_letter(c_sum + 1), get_column_letter(c_sum + 2)
        ws.cell(r, c_sum, f'=IF({mx("長期")}>=-{mn("長期")},{mx("長期")},{mn("長期")})')
        ws.cell(r, c_sum + 1, f'={mx("短期")}')
        ws.cell(r, c_sum + 2, f'={mn("短期")}')
        gov = f'IF({F}{r}>=-{G}{r},{F}{r},{G}{r})'
        ws.cell(r, c_sum + 3, f'=IFERROR(INDEX({name_rng},MATCH(1,INDEX(({cls_rng}="短期")*({rg}={gov}),0),0)),"")')
        ws.cell(r, c_sum + 4, f'={mx("終局")}')
        ws.cell(r, c_sum + 5, f'={mn("終局")}')
        for j in range(len(heads)):
            c = ws.cell(r, 1 + j)
            c.border = BORDER
            if zebra:
                c.fill = ZEBRA
            if c_sum <= 1 + j < c_case:
                c.fill = SUMF
                if j != c_sum - 1 + 3:
                    c.number_format = '0.0'
    for j in range(len(info_heads)):
        ws.column_dimensions[get_column_letter(1 + j)].width = 44 if j == 0 else 22 if j == 1 else 24
    ws.column_dimensions[get_column_letter(c_comp)].width = 12
    ws.column_dimensions[get_column_letter(c_unit)].width = 7
    for j in range(len(sums)):
        ws.column_dimensions[get_column_letter(c_sum + j)].width = 13 if j != 3 else 20
    for j in range(len(cases)):
        ws.column_dimensions[get_column_letter(c_case + j)].width = 13
    ws.freeze_panes = ws.cell(first, c_sum)
    ws.auto_filter.ref = f'A{head_row}:{get_column_letter(last_case)}{first + max(len(rows) - 1, 0)}'
    return dict(first=first, cls_row=cls_row, head_row=head_row, c_sum=c_sum, c_case=c_case, last_case=last_case)


def case_values(results, eid, pt, ci, cases):
    return [results.get(eid, {}).get(c, {}).get(pt, [None] * 6)[ci] for c in cases]


def sheet_columns(ws, model, res, cases, class_cells, columns, face, gi):
    top = title(ws, '4-1 RC柱の設計　設計用応力（midas 梁要素の断面力）', [
        f'対象：RC柱 {len(columns)} 本（断面名が「RC」で始まり、ほぼ鉛直の要素）。値は midas の要素座標系のまま（単位 kN・kN·m、ANL: {res["units"]}）。',
        '柱頭＝上側の節点、柱脚＝下側の節点。フェイス距離を指定した場合は、その距離だけ内側の値を要素の5点から線形補間する。',
        f'フェイス距離：{face:g} mm（0 のとき節点位置の値）。区分は「概要」シートで変えられる（まとめの数式が追従）。',
    ])
    rows, groups = [], []
    for line in columns:
        e = line.elems[0]
        nodes = model['nodes']
        a, b = e['nodes'][0], e['nodes'][1]
        top_is_i = nodes[a][2] >= nodes[b][2]
        L = line.total
        mg = gi.label(nodes[a if top_is_i else b], line.min_z(model)) if gi.available else ''
        for lab, s, nd in (('柱頭', 0 if top_is_i else L, a if top_is_i else b), ('柱中央', L / 2, None), ('柱脚', L if top_is_i else 0, b if top_is_i else a)):
            if lab != '柱中央' and face:
                s = s + face if s == 0 else s - face
            pt3 = model['nodes'][nd] if nd else line.point_at(model, L / 2)
            info = [f'RC柱 {mg}　要素{e["id"]}（{model["sections"].get(e["sec"], "")}） L={L / 1000:.2f}m' if mg else f'RC柱 要素{e["id"]}（{model["sections"].get(e["sec"], "")}） L={L / 1000:.2f}m', f'{lab}' + (f'（節点{nd}）' if nd else ''),
                    gi.label(pt3, line.min_z(model))]
            groups.append((info[0], info[1], info[2], len(rows)))
            for ci, comp in enumerate(COMPS):
                rows.append((info, comp, [line.value_at(res['results'], c, ci, s) for c in cases]))
    meta = write_case_table(ws, top, ['部材', '位置', '通り'], rows, cases, class_cells)
    return meta, [(a, b, g, meta['first'] + k) for a, b, g, k in groups]


def sheet_beams(ws, model, res, cases, class_cells, chains, face, gi):
    top = title(ws, '4-2 RC梁（基礎梁）の設計　設計用応力（midas 梁要素の断面力）', [
        f'対象：RC梁 {len(chains)} 部材（断面名が「RC」で始まる水平な要素。他の梁・支点に接しない節点でつながる要素は 1 部材にまとめた）。',
        f'値は midas の要素座標系のまま（単位 kN・kN·m、ANL: {res["units"]}）。左端・右端＝部材の始端・終端、中央＝部材長の中点（要素の5点から線形補間）。',
        f'フェイス距離：{face:g} mm（0 のとき節点位置の値）。区分は「概要」シートで変えられる（まとめの数式が追従）。',
    ])
    rows, groups = [], []
    for k, line in enumerate(chains, 1):
        L = line.total
        d = min(face, L / 2 - 1) if face else 0
        n0, n1 = line.node_ids[0], line.node_ids[-1]
        mg = gi.member(model['nodes'][n0], model['nodes'][n1], line.min_z(model)) if gi.available else ''
        for lab, s, nd in (('左端', d, n0), ('中央', L / 2, None), ('右端', L - d, n1)):
            pt3 = model['nodes'][nd] if nd else line.point_at(model, L / 2)
            info = [f'RC梁 B{k:02d}　{mg}　要素{id_range(line.elems)}　L={L / 1000:.2f}m' if mg else f'RC梁 B{k:02d}　要素{id_range(line.elems)}　L={L / 1000:.2f}m', lab + (f'（節点{nd}）' if nd else ''),
                    gi.label(pt3, line.min_z(model))]
            groups.append((info[0], info[1], info[2], len(rows)))
            for ci, comp in enumerate(COMPS):
                rows.append((info, comp, [line.value_at(res['results'], c, ci, s) for c in cases]))
    meta = write_case_table(ws, top, ['部材', '位置', '通り'], rows, cases, class_cells)
    return meta, [(a, b, g, meta['first'] + k) for a, b, g, k in groups]


BEAM_PICKS = [
    ('長期', '長期（NΣDL）', 'My', 'long'),
    ('短期', '曲げ My 最大（＋）　下端引張', 'My', 'max'),
    ('短期', '曲げ My 最小（−）　上端引張', 'My', 'min'),
    ('短期', 'せん断 Qz 絶対最大', 'Qz', 'abs'),
    ('終局', '曲げ My 最大（＋）　下端引張', 'My', 'max'),
    ('終局', '曲げ My 最小（−）　上端引張', 'My', 'min'),
    ('終局', 'せん断 Qz 絶対最大', 'Qz', 'abs'),
]
COLUMN_PICKS = [('長期', '長期（NΣDL）', 'N', 'long')]
for _cls in ('短期', '終局'):
    COLUMN_PICKS += [(_cls, '曲げ My 絶対最大', 'My', 'abs'), (_cls, '曲げ Mz 絶対最大', 'Mz', 'abs'),
                     (_cls, 'せん断 Qy 絶対最大', 'Qy', 'abs'), (_cls, 'せん断 Qz 絶対最大', 'Qz', 'abs'),
                     (_cls, '圧縮 最大（N 最小）', 'N', 'min'), (_cls, '引張 最大（N 最大）', 'N', 'max')]


def sheet_chart(ws, heading, notes, detail_name, meta, groups, picks, chart_comps):
    """チャートに入れる値の一覧。詳細シートの数式結果から、同じ荷重ケースの組（N・Qy・Qz・T・My・Mz）を引く。

    選定の基準（最大・最小・絶対最大）で荷重ケースを 1 つ決め、そのケースの 6 成分をすべて並べる。
    チャートに入れる成分は色付き、それ以外は参考（灰色）。
    """
    top = title(ws, heading, notes)
    q = "'" + detail_name + "'"
    cc, lc = get_column_letter(meta['c_case']), get_column_letter(meta['last_case'])
    cls_rng = f"{q}!${cc}${meta['cls_row']}:${lc}${meta['cls_row']}"
    name_rng = f"{q}!${cc}${meta['head_row']}:${lc}${meta['head_row']}"
    sc = lambda off: get_column_letter(meta['c_sum'] + off)   # 0 長期 / 1 短期最大 / 2 短期最小 / 4 終局最大 / 5 終局最小
    heads = ['部材', '位置', '通り', '区分', '選定の基準', '荷重ケース'] + [f'{c}\n({COMP_UNIT[c]})' for c in COMPS] + ['ケース番号']
    for j, h in enumerate(heads):
        c = ws.cell(top, 1 + j, h)
        c.font, c.fill, c.border = BOLD, HEAD, BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[top].height = 34
    r = top + 1
    for gi, (mem, pos, grid, r0) in enumerate(groups):
        for cls, label, drive, mode in picks:
            row_of = lambda comp: r0 + COMPS.index(comp)
            drv = row_of(drive)
            rng = f"{q}!${cc}{drv}:${lc}{drv}"
            if mode == 'long':
                k = f'=IFERROR(MATCH("長期",{cls_rng},0),"")'
            else:
                mx, mn = (f"{q}!{sc(1 if cls == '短期' else 4)}{drv}", f"{q}!{sc(2 if cls == '短期' else 5)}{drv}")
                val = mx if mode == 'max' else mn if mode == 'min' else f'IF({mx}>=-{mn},{mx},{mn})'
                k = f'=IFERROR(MATCH(1,INDEX(({cls_rng}="{cls}")*({rng}={val}),0),0),"")'
            cells = [mem, pos, grid, cls, label, f'=IF($M{r}="","",INDEX({name_rng},$M{r}))']
            for comp in COMPS:
                rr = row_of(comp)
                cells.append(f'=IF($M{r}="","",INDEX({q}!${cc}{rr}:${lc}{rr},$M{r}))')
            cells.append(k)
            for j, v in enumerate(cells):
                c = ws.cell(r, 1 + j, v)
                c.border = BORDER
                if 6 <= j <= 11:
                    c.number_format = '0.0'
                    comp = COMPS[j - 6]
                    if comp in chart_comps:
                        c.fill = SUMF
                        c.font = BOLD
                    else:
                        c.font = Font(color='999999')
                if j == 12:
                    c.font = Font(color='999999', size=8)
                if cls == '長期' and j < 6:
                    c.fill = PatternFill('solid', fgColor=CLASS_COLOR['長期'])
                elif j == 3:
                    c.fill = PatternFill('solid', fgColor=CLASS_COLOR[cls])
            r += 1
    for j, w in enumerate((44, 22, 24, 7, 30, 18, 10, 10, 10, 10, 10, 10, 8)):
        ws.column_dimensions[get_column_letter(1 + j)].width = w
    ws.freeze_panes = ws.cell(top + 1, 7)
    ws.auto_filter.ref = f'A{top}:M{max(r - 1, top)}'
    return dict(first=top + 1, n=len(picks))


def sheet_grids(ws, model, gi):
    top = title(ws, '通り芯（midas の寸法線・テキストから読み取り）', [
        '表記はチャートの「位置」欄と同じ。点：「BX2・BY1」＝通り上（100 mm 以内）、「BX2+1250・BY2」＝最寄りの通りから 1250 mm（＋は番号が進む側＝X は東、Y は北）。',
        '部材：通りに沿うものは「BY1=BX1～BX2」（通り＝始点～終点）、沿わないもの（円弧など）は「始点 ～ 終点」。',
        '面の選び方：RC部材・支点まわりは、部材の最低点の高さに一番近い寸法ビュー（基礎レベルの BX・BY）、鉄骨は屋根レベルの X・Y を使う。座標は m（midas の全体座標）。',
    ])
    for j, h in enumerate(['寸法ビュー', '符号', '系', '始点 X', '始点 Y', '終点 X', '終点 Y', '高さ Z', '角度(°)']):
        c = ws.cell(top, 1 + j, h)
        c.font, c.fill, c.border = BOLD, HEAD, BORDER
    r = top + 1
    for v in gi.views.values():
        for fam in ('X', 'Y'):
            for g in v['fam'].get(fam, []):
                ang = math.degrees(math.atan2(g['p1'][1] - g['p0'][1], g['p1'][0] - g['p0'][0]))
                vals = [g['view'], g['name'], fam, g['p0'][0] / 1000, g['p0'][1] / 1000, g['p1'][0] / 1000, g['p1'][1] / 1000, g['z'] / 1000, ang]
                for j, x in enumerate(vals):
                    c = ws.cell(r, 1 + j, round(x, 4) if isinstance(x, float) else x)
                    c.border = BORDER
                    if isinstance(x, float):
                        c.number_format = '0.000'
                r += 1
    for j, w in enumerate((14, 8, 5, 12, 12, 12, 12, 9, 9)):
        ws.column_dimensions[get_column_letter(1 + j)].width = w


def sheet_lcr(ws, chart_name, chart, groups, picks, positions):
    """部材ごとに 左端・中央・右端 を横に並べた一覧。値はチャート入力シートのセルを参照する。

    曲げ（My）とせん断（Qz）を別の行にして、長期・短期・終局の順に並べる（計算書の「応力算定」と同じ並び）。
    """
    q = "'" + chart_name + "'"
    top = title(ws, '4-2 RC梁　左端・中央・右端の設計用応力（チャート入力）', [
        '値はチャート入力シート（' + chart_name + '）の同じ値を参照している。曲げは My（＋＝下端引張、−＝上端引張）、せん断は Qz。単位は kN・kN·m。',
        '右の「ケース」欄は、その値が出た荷重ケース（位置ごと）。区分・選定の基準はチャート入力シートと同じ。',
    ])
    heads = ['部材（通り）', '区分', '項目', '単位'] + positions + [f'{p} ケース' for p in positions]
    for j, h in enumerate(heads):
        c = ws.cell(top, 1 + j, h)
        c.font, c.fill, c.border = BOLD, HEAD, BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    items = [('長期', '曲げ My', 'K', 0), ('長期', 'せん断 Qz', 'I', 0)]
    names = {1: '曲げ My 最大（＋）　下端引張', 2: '曲げ My 最小（−）　上端引張', 3: 'せん断 Qz 絶対最大'}
    for cls, base in (('短期', 1), ('終局', 4)):
        items += [(cls, names[1], 'K', base), (cls, names[2], 'K', base + 1), (cls, names[3], 'I', base + 2)]
    npos = len(positions)
    r = top + 1
    for mi in range(len(groups) // npos):
        label = groups[mi * npos][0]
        for cls, item, col, pi in items:
            ws.cell(r, 1, label)
            ws.cell(r, 2, cls).fill = PatternFill('solid', fgColor=CLASS_COLOR[cls])
            ws.cell(r, 3, item)
            ws.cell(r, 4, 'kN·m' if col == 'K' else 'kN')
            for p in range(npos):
                cr = chart['first'] + (mi * npos + p) * chart['n'] + pi
                v = ws.cell(r, 5 + p, f'={q}!{col}{cr}')
                v.number_format = '0.0'
                v.font = BOLD
                v.fill = SUMF
                k = ws.cell(r, 5 + npos + p, f'={q}!F{cr}')
                k.font = Font(color='666666', size=9)
            for j in range(len(heads)):
                ws.cell(r, 1 + j).border = BORDER
            r += 1
    for j, w in enumerate([66, 7, 32, 7] + [11] * npos + [18] * npos):
        ws.column_dimensions[get_column_letter(1 + j)].width = w
    ws.freeze_panes = ws.cell(top + 1, 5)
    ws.auto_filter.ref = f'A{top}:{get_column_letter(len(heads))}{max(r - 1, top)}'


def sheet_walls(ws, model):
    plates = [e for e in model['elements'].values() if e['type'] == 'PLATE']
    top = title(ws, '4-3 耐震壁の設計　設計用応力', [
        '耐震壁（板要素）の応力は、今回の ANL には出力されていない（ANL は梁要素の断面力のみ）。',
        'midas で「板要素の断面力」を ANL に出力すると、壁の応力も抜き出せる（出力できたらサンプルを渡してもらえれば読み込みを追加する）。',
        '参考として、モデルにある板要素の一覧を示す。',
    ])
    for j, h in enumerate(['要素', '厚さ t (mm)', '節点']):
        c = ws.cell(top, 1 + j, h)
        c.font, c.fill, c.border = BOLD, HEAD, BORDER
    for i, e in enumerate(plates):
        r = top + 1 + i
        ws.cell(r, 1, e['id'])
        t = model['thickness'].get(e['sec'])
        ws.cell(r, 2, t)
        ws.cell(r, 3, ' '.join(str(n) for n in e['nodes']))
        for j in range(3):
            ws.cell(r, 1 + j).border = BORDER
    for col, w in zip('ABC', (10, 14, 40)):
        ws.column_dimensions[col].width = w


def sheet_supports(ws, model, res, cases, class_cells, gi):
    results = res['results']
    top = title(ws, '4-4 柱脚・支点まわり　支点節点に接する要素の端部力', [
        f'支点（拘束）節点に接する梁要素について、支点側の端部の断面力を並べた。柱脚・アンカー・接合部の検討に使う（単位 kN・kN·m、ANL: {res["units"]}）。',
        '値は midas の要素座標系のまま。支点反力そのもの（全体座標）ではない。反力表は midas の反力出力（6-2 支点反力表）を使うこと。',
    ])
    rows = []
    for n in sorted(model['supports']):
        for e in sorted(model['elements'].values(), key=lambda x: x['id']):
            if e['type'] != 'BEAM' or n not in e['nodes'][:2]:
                continue
            end = 'I' if e['nodes'][0] == n else 'J'
            info = [f'節点 {n}　要素{e["id"]}（{model["sections"].get(e["sec"], "")}）', f'{end}端', gi.label(model['nodes'][n], model['nodes'][n][2])]
            for ci, comp in enumerate(COMPS):
                rows.append((info, comp, case_values(results, e['id'], end, ci, cases)))
    write_case_table(ws, top, ['支点節点・要素（断面）', '位置', '通り'], rows, cases, class_cells)


def extremes(elem_ids, results, classes):
    """要素群の中で、成分×区分ごとの最大・最小とその位置（要素・点・ケース）を求める。"""
    best = {}
    for eid in elem_ids:
        for case, pts in results.get(eid, {}).items():
            cls = classes.get(case)
            if not cls:
                continue
            for pt, vals in pts.items():
                for ci, v in enumerate(vals):
                    for kind in ('max', 'min'):
                        key = (ci, cls, kind)
                        cur = best.get(key)
                        if cur is None or (v > cur[0] if kind == 'max' else v < cur[0]):
                            best[key] = (v, eid, pt, case)
    return best


def sheet_summary(ws, heading, note, groups, model, res, classes, gi):
    top = title(ws, heading, [note, f'単位 kN・kN·m（ANL: {res["units"]}）。区分は生成時の仮決め（長期＝NΣDL、短期＝SL・SLU・WZ・EX・EY、終局＝（終局））。位置は「要素 点　通り　ケース」。'])
    heads = ['名称', '要素数', '成分', '単位', '長期 絶対最大', '長期 位置', '短期 最大', '短期 位置', '短期 最小', '短期 位置', '終局 最大', '終局 位置', '終局 最小', '終局 位置']
    for j, h in enumerate(heads):
        c = ws.cell(top, 1 + j, h)
        c.font, c.fill, c.border = BOLD, HEAD, BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    r = top + 1
    zebra = False
    for name, elem_ids in groups:
        ids = [i for i in elem_ids if i in res['results']]
        if not ids:
            continue
        best = extremes(ids, res['results'], classes)
        zebra = not zebra
        for ci, comp in enumerate(COMPS):
            def pick(cls, kind):
                v = best.get((ci, cls, kind))
                if not v:
                    return (None, '')
                grid = element_grid(model, gi, v[1], v[2]) if gi.available else ''
                return (round(v[0], 3), f'{v[1]} {v[2]}　{grid}　{v[3]}' if grid else f'{v[1]} {v[2]}　{v[3]}')
            lo, hi = pick('長期', 'max'), pick('長期', 'min')
            long_abs = lo if lo[0] is not None and (hi[0] is None or lo[0] >= -hi[0]) else hi
            cells = [name if ci == 0 else '', len(ids) if ci == 0 else '', COMP_NAME[comp], COMP_UNIT[comp],
                     long_abs[0], long_abs[1], *pick('短期', 'max'), *pick('短期', 'min'), *pick('終局', 'max'), *pick('終局', 'min')]
            for j, v in enumerate(cells):
                c = ws.cell(r, 1 + j, v)
                c.border = BORDER
                if zebra:
                    c.fill = ZEBRA
                if j in (4, 6, 8, 10, 12):
                    c.number_format = '0.0'
            r += 1
    for j, w in enumerate((30, 8, 12, 7, 12, 44, 12, 44, 12, 44, 12, 44, 12, 44)):
        ws.column_dimensions[get_column_letter(1 + j)].width = w
    ws.freeze_panes = ws.cell(top + 1, 4)
    ws.auto_filter.ref = f'A{top}:N{max(r - 1, top)}'


def sheet_overview(ws, model_path, anl_path, model, res, cases, classes, n_col, n_beam, face_beam, face_col):
    ws['A1'] = '日吉津グランルーフ　midas 応力の抜き出し'
    ws['A1'].font = Font(bold=True, size=14)
    info = [('モデル', os.path.basename(model_path)), ('解析結果', os.path.basename(anl_path)),
            ('単位', f'ANL は {res["units"]} → 表は kN・kN·m に換算'),
            ('要素数', f'梁 {sum(1 for e in model["elements"].values() if e["type"] == "BEAM")}・板 {sum(1 for e in model["elements"].values() if e["type"] == "PLATE")}（ANL に断面力のある梁要素 {len(res["results"])}）'),
            ('RC部材', f'RC柱 {n_col} 本・RC梁 {n_beam} 部材（フェイス距離：梁 {face_beam:g} mm／柱 {face_col:g} mm）')]
    for i, (k, v) in enumerate(info):
        ws.cell(3 + i, 1, k).font = BOLD
        ws.cell(3 + i, 2, v)
    r0 = 10
    ws.cell(r0 - 1, 1, '荷重ケースと区分（黄色のセルは変更できる。RC のシートのまとめ数式が追従する）').font = BOLD
    for j, h in enumerate(['No.', '荷重ケース', '区分']):
        c = ws.cell(r0, 1 + j, h)
        c.font, c.fill, c.border = BOLD, HEAD, BORDER
    refs = []
    dv = DataValidation(type='list', formula1='"長期,短期,終局,対象外"', allow_blank=False)
    ws.add_data_validation(dv)
    for i, name in enumerate(cases):
        r = r0 + 1 + i
        ws.cell(r, 1, i + 1)
        ws.cell(r, 2, name)
        c = ws.cell(r, 3, classes[name])
        c.fill = INPF
        dv.add(c)
        for j in range(3):
            ws.cell(r, 1 + j).border = BORDER
        refs.append(f"'概要'!$C${r}")
    r = r0 + len(cases) + 3
    ws.cell(r, 1, 'シート').font = BOLD
    guide = [('4-1_RC柱', '【チャート入力】RC柱の柱頭・柱脚で、チャートに入れる値（色付き）を荷重ケースの組で並べた表'),
             ('4-1_RC柱(詳細)', 'RC柱の全荷重ケースの応力と、長期／短期／終局のまとめ（数式）'),
             ('4-2_RC梁', '【チャート入力】RC梁（基礎梁）の左端・中央・右端で、チャートに入れる値（色付き）'),
             ('4-2_RC梁(左中右)', '【左端・中央・右端】部材ごとに、曲げ・せん断を左端／中央／右端の横並びにした一覧（計算書の応力算定と同じ並び）'),
             ('4-2_RC梁(詳細)', 'RC梁の全荷重ケースの応力と、長期／短期／終局のまとめ（数式）'),
             ('4-3_耐震壁', '壁（板要素）の応力は ANL に無いため一覧のみ'),
             ('4-4_支点まわり', '支点節点に接する要素の端部力（柱脚・接合部の検討用）'),
             ('通り芯', '「通り」列の判定に使った、midas の通り芯の一覧'),
             ('断面別まとめ', '断面ごとの最大・最小（長期／短期／終局）と、その要素・点・ケース。鉄骨部材の検定用'),
             ('グループ別まとめ', 'midas のグループ（トラス・つなぎ・キール・リング）ごとの最大・最小')]
    for i, (k, v) in enumerate(guide):
        ws.cell(r + 1 + i, 1, k)
        ws.cell(r + 1 + i, 2, v)
    r += len(guide) + 3
    ws.cell(r, 1, '注意').font = BOLD
    notes = ['ANL に出ているのは梁要素の断面力のみ。板要素（耐震壁）の応力と支点反力（6-2・基礎）は出力されていないので、この Excel には入らない。',
             '値は midas の要素座標系の断面力（要素ごとの向き）。符号の約束は midas に従う。',
             '短期の「最大・最小」は符号付き。支配ケース＝短期の絶対値が最大になるケース。',
             'ANL の荷重ケース名は表ごとに短縮されているため、要約表の正式名称に置き換えている（例 NΣDL+S~4 → NΣDL+SLU1）。',
             '地震時（EX・EY）は短期に含めた。必要保有水平耐力時の検討は「終局」を使う。']
    for i, n in enumerate(notes):
        ws.cell(r + 1 + i, 1, '・' + n)
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 10
    return refs


def build(model_path, anl_path, out_path, face_beam=0.0, face_col=0.0, log=say, model_label=None, anl_label=None):
    log('モデルを読み込み中…')
    model = read_mgt(model_path)
    log(f'  節点 {len(model["nodes"])}・要素 {len(model["elements"])}・グループ {len(model["groups"])}・支点節点 {len(model["supports"])}')
    log('解析結果を読み込み中…（大きいと少し時間がかかります）')
    res = read_anl(anl_path)
    cases = res['cases']
    log(f'  荷重ケース {len(cases)}・梁要素 {len(res["results"])}：' + ' / '.join(cases))
    if not res['results']:
        raise ValueError('ANL から梁要素の断面力を読み取れませんでした（梁要素の断面力を出力した ANL か確認してください）。')
    classes = {c: classify(c) for c in cases}

    rc_secs = {i for i, n in model['sections'].items() if n.upper().startswith('RC')}
    rc = [e for e in model['elements'].values() if e['type'] == 'BEAM' and e['sec'] in rc_secs]

    def vertical(e):
        a, b = (model['nodes'][n] for n in e['nodes'][:2])
        return abs(b[2] - a[2]) > 0.7 * math.dist(a, b)
    col_elems = [e for e in rc if vertical(e)]
    beam_elems = [e for e in rc if not vertical(e)]
    columns = [Line([e], model) for e in sorted(col_elems, key=lambda x: x['id'])]
    chains = build_chains(beam_elems, model)
    log(f'  RC柱 {len(columns)} 本・RC梁 {len(chains)} 部材（{len(beam_elems)} 要素）')

    wb = Workbook()
    ws0 = wb.active
    ws0.title = '概要'
    gi = GridIndex(model['grids'])
    log(f'  通り芯 {len(model["grids"])} 本（寸法ビュー：{", ".join(gi.views) or "なし"}）')
    refs = sheet_overview(ws0, model_label or model_path, anl_label or anl_path, model, res, cases, classes, len(columns), len(chains), face_beam, face_col)
    ws_c1 = wb.create_sheet('4-1_RC柱')
    meta, groups = sheet_columns(wb.create_sheet('4-1_RC柱(詳細)'), model, res, cases, refs, columns, face_col, gi)
    sheet_chart(ws_c1, '4-1 RC柱　チャートに入れる設計用応力', [
        '色付きのセルがチャート（軸力＋曲げ／せん断）に入れる値。同じ行は同じ荷重ケースの組なので、N・My・Mz はそのまま組で使う。',
        '位置は柱頭・柱脚（RC部材の検定位置）。短期・終局は、選定の基準（曲げ最大・せん断最大・圧縮最大・引張最大）ごとに 1 ケースを選んでいる。',
        'N は midas の符号で、圧縮が負（長期で負になる）。チャートに入れるときは圧縮を正に読み替える。My・Mz は絶対値で使う（2方向曲げは別々に検定）。',
        '短期の各ケースは、すでに長期荷重を含む組合せ（NΣDL＋…）なので、長期との和は取らない。T は参考。全ケースの値は「4-1_RC柱(詳細)」。',
    ], '4-1_RC柱(詳細)', meta, groups, COLUMN_PICKS, {'N', 'Qy', 'Qz', 'My', 'Mz'})
    ws_b1 = wb.create_sheet('4-2_RC梁')
    ws_lcr = wb.create_sheet('4-2_RC梁(左中右)')
    meta, groups = sheet_beams(wb.create_sheet('4-2_RC梁(詳細)'), model, res, cases, refs, chains, face_beam, gi)
    chart_b = sheet_chart(ws_b1, '4-2 RC梁（基礎梁）　チャートに入れる設計用応力', [
        '色付きのセルがチャート（曲げ・せん断）に入れる値。曲げ：下端筋は My が＋の最大、上端筋は My が−の最大（絶対値）。せん断：Qz の絶対最大。',
        '位置は左端・中央・右端。midas の My は、スパン中央で＋（下端引張）、端部で−（上端引張）になる符号。N・T・Mz・Qy は参考（灰色）。',
        '短期の各ケースは、すでに長期荷重を含む組合せ（NΣDL＋…）なので、長期との和は取らない。終局は 5-3 保有時の断面算定用。',
        '全ケースの値と、長期／短期／終局のまとめは「4-2_RC梁(詳細)」。',
    ], '4-2_RC梁(詳細)', meta, groups, BEAM_PICKS, {'My', 'Qz'})
    sheet_lcr(ws_lcr, '4-2_RC梁', chart_b, groups, BEAM_PICKS, ['左端', '中央', '右端'])
    sheet_walls(wb.create_sheet('4-3_耐震壁'), model)
    sheet_supports(wb.create_sheet('4-4_支点まわり'), model, res, cases, refs, gi)

    by_sec = collections.OrderedDict()
    for e in sorted(model['elements'].values(), key=lambda x: x['id']):
        if e['type'] == 'BEAM' and e['id'] in res['results']:
            by_sec.setdefault(e['sec'], []).append(e['id'])
    sec_groups = [(f'{s}  {model["sections"].get(s, "")}', ids) for s, ids in sorted(by_sec.items())]
    sheet_summary(wb.create_sheet('断面別まとめ'), '断面ごとの最大・最小応力（全梁要素）',
                  '各断面の全要素・全点・全荷重ケースの中での最大・最小。トラス・つなぎ・キール・リング各部材の検定に使う。', sec_groups, model, res, classes, gi)
    grp = [(n, ids) for n, ids in model['groups'].items() if ids and not n.startswith('膜')]
    sheet_summary(wb.create_sheet('グループ別まとめ'), 'midas グループごとの最大・最小応力',
                  'グループ（トラス_X1〜X5、つなぎ_Ya〜Yn、キール、リング）に入っている要素の中での最大・最小。', grp, model, res, classes, gi)
    if gi.available:
        sheet_grids(wb.create_sheet('通り芯'), model, gi)
    wb.save(out_path)
    log(f'保存しました：{out_path}')


# ───────────────────────── 画面 ─────────────────────────

def run_gui():
    import queue
    import threading
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title('midas 応力抽出 Excel')
    root.geometry('720x520')
    v_model, v_anl, v_out = tk.StringVar(), tk.StringVar(), tk.StringVar()
    v_fb, v_fc = tk.StringVar(value='0'), tk.StringVar(value='0')
    q = queue.Queue()

    def pick(var, kinds, save=False):
        p = (filedialog.asksaveasfilename(defaultextension='.xlsx', filetypes=kinds) if save else filedialog.askopenfilename(filetypes=kinds))
        if p:
            var.set(p)
            if var is v_model and not v_out.get():
                v_out.set(os.path.splitext(p)[0] + '_応力抽出.xlsx')

    def row(r, label, var, kinds, save=False):
        tk.Label(root, text=label, anchor='w').grid(row=r, column=0, sticky='w', padx=8, pady=4)
        tk.Entry(root, textvariable=var, width=62).grid(row=r, column=1, padx=4)
        tk.Button(root, text='選ぶ', command=lambda: pick(var, kinds, save)).grid(row=r, column=2, padx=6)

    row(0, 'モデル（MGT/MGTX）', v_model, [('MGT/MGTX', '*.mgt *.mgtx'), ('すべて', '*.*')])
    row(1, '解析結果（ANL）', v_anl, [('ANL', '*.anl'), ('すべて', '*.*')])
    row(2, '出力 Excel', v_out, [('Excel', '*.xlsx')], save=True)
    tk.Label(root, text='フェイス距離 mm（0＝節点位置）  RC梁:').grid(row=3, column=0, columnspan=2, sticky='w', padx=8)
    tk.Entry(root, textvariable=v_fb, width=8).grid(row=3, column=1, sticky='e', padx=(0, 120))
    tk.Label(root, text='RC柱:').grid(row=3, column=1, sticky='e', padx=(0, 80))
    tk.Entry(root, textvariable=v_fc, width=8).grid(row=3, column=1, sticky='e')
    log_box = scrolledtext.ScrolledText(root, height=20, width=90)
    log_box.grid(row=5, column=0, columnspan=3, padx=8, pady=8)
    btn = tk.Button(root, text='Excel を作る', width=20, height=2)
    btn.grid(row=4, column=1, pady=6)

    def work():
        try:
            build(v_model.get(), v_anl.get(), v_out.get(), float(v_fb.get() or 0), float(v_fc.get() or 0), log=lambda m: q.put(m))
            q.put('__done__')
        except BaseException as exc:  # noqa
            q.put(f'エラー：{exc}')
            q.put('__done__')

    def start():
        if not (os.path.isfile(v_model.get()) and os.path.isfile(v_anl.get()) and v_out.get()):
            messagebox.showwarning('確認', 'モデル・解析結果・出力先を指定してください。')
            return
        btn.config(state='disabled')
        threading.Thread(target=work, daemon=True).start()

    def poll():
        try:
            while True:
                m = q.get_nowait()
                if m == '__done__':
                    btn.config(state='normal')
                else:
                    log_box.insert('end', m + '\n')
                    log_box.see('end')
        except queue.Empty:
            pass
        root.after(150, poll)

    btn.config(command=start)
    poll()
    root.mainloop()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('model', nargs='?')
    ap.add_argument('anl', nargs='?')
    ap.add_argument('-o', '--out')
    ap.add_argument('--face-beam', type=float, default=0.0, help='RC梁のフェイス距離 mm')
    ap.add_argument('--face-col', type=float, default=0.0, help='RC柱のフェイス距離 mm')
    a = ap.parse_args()
    if not a.model:
        return run_gui()
    if not a.anl:
        ap.error('ANL ファイルも指定してください')
    out = a.out or os.path.splitext(a.model)[0] + '_応力抽出.xlsx'
    try:
        build(a.model, a.anl, out, a.face_beam, a.face_col)
    except ValueError as exc:
        sys.exit(str(exc))


if __name__ == '__main__':
    main()
